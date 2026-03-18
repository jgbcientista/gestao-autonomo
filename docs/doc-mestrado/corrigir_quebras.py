import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

DOCX_PATH = 'doc-mestrado/doc-qualificacao-ifba.docx'
doc = Document(DOCX_PATH)

# Strategy: add pageBreakBefore to each chapter BOX paragraph
# This ensures each chapter starts on a new page

print("=== Adicionando quebra de página antes de cada capítulo ===\n")

for i, p in enumerate(doc.paragraphs):
    xml_str = etree.tostring(p._element, pretty_print=True).decode()
    if 'w:drawing' not in xml_str:
        continue

    texts = [t.text for t in p._element.iter(qn('w:t')) if t.text and t.text.strip()]
    combined = ' '.join(texts)
    if 'CAPÍTULO' not in combined:
        continue

    # Find next H1
    next_h1 = ''
    for j in range(i+1, min(i+15, len(doc.paragraphs))):
        if doc.paragraphs[j].style and doc.paragraphs[j].style.name == 'Heading 1' and doc.paragraphs[j].text.strip():
            next_h1 = doc.paragraphs[j].text[:50]
            break

    # Check if there's already a section break before this box
    # Look at the 1-3 paragraphs before the box
    has_break = False
    for k in range(max(0, i-3), i):
        pk = doc.paragraphs[k]
        pkPr = pk._element.find(qn('w:pPr'))
        if pkPr is not None:
            sectPr = pkPr.find(qn('w:sectPr'))
            if sectPr is not None:
                has_break = True
                break
        # Also check for pageBreakBefore
        if pkPr is not None:
            pgBrk = pkPr.find(qn('w:pageBreakBefore'))
            if pgBrk is not None:
                has_break = True
                break

    # Also check the box paragraph itself
    pPr = p._element.find(qn('w:pPr'))
    if pPr is not None:
        if pPr.find(qn('w:pageBreakBefore')) is not None:
            has_break = True

    if has_break:
        print(f"  {combined} → {next_h1} (já tem quebra)")
    else:
        # Add pageBreakBefore to the box paragraph
        pPr = p._element.find(qn('w:pPr'))
        if pPr is None:
            pPr = OxmlElement('w:pPr')
            p._element.insert(0, pPr)

        pgBrk = OxmlElement('w:pageBreakBefore')
        pPr.insert(0, pgBrk)
        print(f"  {combined} → {next_h1} (quebra adicionada)")

# Also ensure MOTOR H2 subcapítulos DON'T have unwanted page breaks
print("\n=== Verificando subcapítulos MOTOR (sem quebras indevidas) ===")
motor_start = None
dev_start = None
for i, p in enumerate(doc.paragraphs):
    if p.style and p.style.name == 'Heading 1':
        if 'MOTOR' in p.text:
            motor_start = i
        elif 'DESENVOLVIMENTO' in p.text:
            dev_start = i
            break

if motor_start and dev_start:
    for i in range(motor_start, dev_start):
        p = doc.paragraphs[i]
        if p.style and p.style.name == 'Heading 2':
            pPr = p._element.find(qn('w:pPr'))
            if pPr is not None:
                pgBrk = pPr.find(qn('w:pageBreakBefore'))
                if pgBrk is not None:
                    pPr.remove(pgBrk)
                    print(f"  {i}: removida quebra indevida de {p.text[:50]}")

doc.save(DOCX_PATH)
print(f"\nSalvo: {DOCX_PATH}")
