# -*- coding: utf-8 -*-
"""Extrair o conteúdo atual do cronograma."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc = Document(r'C:\dev\mestrado\qualificacao-ifba-2025-FINAL.docx')

# Encontrar seção CRONOGRAMA
print("=== SEÇÃO CRONOGRAMA ===")
in_crono = False
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    s = p.style.name if p.style else ''
    if 'CRONOGRAMA' in t and 'Heading' in s:
        in_crono = True
        print(f"\n[{i}] {s}: {t}")
        continue
    if in_crono and 'Heading 1' in s and t and 'CRONOGRAMA' not in t:
        print(f"\n--- FIM (próximo: {t}) ---")
        break
    if in_crono and t:
        print(f"  [{i}] {s}: {t}")

# Tabela do cronograma
print("\n\n=== TABELA DO CRONOGRAMA ===")
for idx, table in enumerate(doc.tables):
    print(f"\nTabela {idx+1}: {len(table.rows)} linhas x {len(table.columns)} colunas")
    for j, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        print(f"  Linha {j}: {cells}")
