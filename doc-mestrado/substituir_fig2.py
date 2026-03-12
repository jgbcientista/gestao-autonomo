import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Emu
from lxml import etree

DOCX_PATH = 'doc-mestrado/doc-qualificacao-ifba.docx'
doc = Document(DOCX_PATH)
part = doc.part

# Find Figura 2 image (para 426)
print("=== Localizando Figura 2 ===")
fig2_para = None
fig2_idx = None

for i, p in enumerate(doc.paragraphs):
    # Check next paragraph for "Figura 2" caption
    if i + 1 < len(doc.paragraphs):
        next_text = doc.paragraphs[i+1].text if doc.paragraphs[i+1].text else ''
        if 'Figura 2' in next_text:
            xml_str = etree.tostring(p._element, pretty_print=True).decode()
            if 'w:drawing' in xml_str:
                fig2_para = p
                fig2_idx = i
                print(f"  Encontrada em para {i}")
                break

if fig2_para is None:
    print("  Figura 2 não encontrada!")
    sys.exit(1)

# Find the image relationship ID
blip = None
for b in fig2_para._element.iter(qn('a:blip')):
    blip = b
    break

if blip is None:
    print("  Blip não encontrado!")
    sys.exit(1)

embed_rid = blip.get(qn('r:embed'))
print(f"  Image r:id = {embed_rid}")

# Get the image part and replace its content
rel = part.rels[embed_rid]
image_part = rel.target_part
print(f"  Content type: {image_part.content_type}")

# Read the new image
with open('diagrama_sequencia.png', 'rb') as f:
    new_image_data = f.read()

# Replace the image data
image_part._blob = new_image_data
print(f"  Imagem substituída ({len(new_image_data)} bytes)")

# Update the image dimensions in the document to fit within margins
# Usable width = 16cm, maintain aspect ratio
# New image: 6.3in x 8.5in = 16cm x 21.6cm -> too tall, scale to fit width
target_width_cm = 15.5  # slightly less than 16cm for safety
target_width_emu = int(target_width_cm / 2.54 * 914400)

# Get current dimensions and calculate aspect ratio from new image
from PIL import Image
img = Image.open('diagrama_sequencia.png')
img_w, img_h = img.size
aspect = img_h / img_w
target_height_emu = int(target_width_emu * aspect)

print(f"  Nova dimensão: {target_width_cm:.1f} x {target_width_cm * aspect:.1f} cm")

# Update extent (wp:extent)
for extent in fig2_para._element.iter(qn('wp:extent')):
    old_cx = extent.get('cx')
    old_cy = extent.get('cy')
    extent.set('cx', str(target_width_emu))
    extent.set('cy', str(target_height_emu))
    print(f"  wp:extent: {old_cx}x{old_cy} → {target_width_emu}x{target_height_emu}")

# Also update a:ext in a:xfrm
for ext in fig2_para._element.iter(qn('a:ext')):
    cx = ext.get('cx')
    cy = ext.get('cy')
    if cx and int(cx) > 100000:  # Only update the main image extent
        ext.set('cx', str(target_width_emu))
        ext.set('cy', str(target_height_emu))
        print(f"  a:ext: {cx}x{cy} → {target_width_emu}x{target_height_emu}")

doc.save(DOCX_PATH)
print(f"\nSalvo: {DOCX_PATH}")
