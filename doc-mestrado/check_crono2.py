# -*- coding: utf-8 -*-
"""Verificar tabelas e conteúdo completo do cronograma."""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

doc = Document(r'C:\dev\mestrado\qualificacao-ifba-2025-FINAL.docx')

# Listar TODOS os parágrafos da seção cronograma até referências
print("=== TODOS OS PARÁGRAFOS DO CRONOGRAMA ===")
in_crono = False
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    s = p.style.name if p.style else ''
    if 'CRONOGRAMA' in t and 'Heading 1' in s:
        in_crono = True
    if in_crono and 'REFER' in t and 'Heading 1' in s:
        break
    if in_crono:
        print(f"  [{i}] {s}: '{t[:80]}'")

# Verificar tabelas
print(f"\n=== TABELAS NO DOCUMENTO: {len(doc.tables)} ===")
for idx, table in enumerate(doc.tables):
    print(f"\nTabela {idx+1}: {len(table.rows)} linhas x {len(table.columns)} colunas")
    for j, row in enumerate(table.rows):
        cells = []
        for cell in row.cells:
            cells.append(cell.text.strip()[:15])
        print(f"  Linha {j}: {cells}")

# Verificar se há tabelas via XML
print("\n=== TABELAS VIA XML ===")
body = doc.element.body
tables = body.findall(qn('w:tbl'))
print(f"Tabelas encontradas via XML: {len(tables)}")
for idx, tbl in enumerate(tables):
    rows = tbl.findall(qn('w:tr'))
    print(f"  Tabela {idx+1}: {len(rows)} linhas")
    for j, row in enumerate(rows):
        cells = row.findall(qn('w:tc'))
        cell_texts = []
        for cell in cells:
            text = ''
            for p in cell.findall(qn('w:p')):
                for r in p.findall(qn('w:r')):
                    for t in r.findall(qn('w:t')):
                        text += t.text or ''
            cell_texts.append(text.strip()[:15])
        print(f"    Linha {j}: {cell_texts}")
