# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Usar o template IFBA como base
prs = Presentation(r"C:\dev\projetos\sistema-autonomo\docs\template-ifba.pptx")

# Remover slides existentes do template
while len(prs.slides) > 0:
    rId = prs.slides._sldIdLst[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])

# Dimensões: 10.00 x 7.50 in (4:3)
SW = prs.slide_width   # 10 in
SH = prs.slide_height  # 7.5 in

# === CORES (padrão PPGESP/IFBA) ===
CINZA_BARRA = RGBColor(0x3A, 0x3A, 0x3A)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
PRETO = RGBColor(0x00, 0x00, 0x00)
CINZA_CLARO = RGBColor(0xF0, 0xF2, 0xF5)
CINZA_TEXTO = RGBColor(0x33, 0x33, 0x33)
VERDE_IF = RGBColor(0x2F, 0x9E, 0x41)
VERDE_ESCURO = RGBColor(0x1A, 0x6B, 0x2A)
VERMELHO_IF = RGBColor(0xCD, 0x19, 0x1E)
LARANJA = RGBColor(0xE6, 0x7E, 0x22)

# Caminhos das imagens extraídas do template
LOGO_PPGESP = r"C:\dev\projetos\sistema-autonomo\docs\logo-ifba.png"
LOGO_CUBO = r"C:\dev\projetos\sistema-autonomo\docs\master-img-Picture 1.png"
BG_CAPA = r"C:\dev\projetos\sistema-autonomo\docs\layout0-10.0x7.5.png"

# === FUNÇÕES AUXILIARES ===

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, color, corner_radius=None):
    if corner_radius:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.adjustments[0] = corner_radius
    else:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_text(slide, left, top, width, height, text, font_size=18, color=CINZA_TEXTO,
             bold=False, alignment=PP_ALIGN.JUSTIFY, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_bullet_list(slide, left, top, width, height, items, font_size=14,
                    color=CINZA_TEXTO, bullet_char="\u2022"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"{bullet_char} {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(4)
        p.alignment = PP_ALIGN.JUSTIFY
    return txBox

def slide_padrao(slide, titulo, subtitulo=None):
    """Aplica o layout padrão dos slides internos do PPGESP"""
    set_slide_bg(slide, BRANCO)
    # Logo cubo no canto superior direito
    slide.shapes.add_picture(LOGO_CUBO, Inches(9.02), Inches(0.21), Inches(0.70), Inches(0.80))
    # Barra inferior cinza escuro
    add_shape(slide, Inches(0), Inches(7.09), Inches(10), Inches(0.31), CINZA_BARRA)
    # Rodapé
    add_text(slide, Inches(0.3), Inches(7.10), Inches(7), Inches(0.25),
             "IFBA – PPGESP – João Guedes de Brito",
             font_size=8, color=BRANCO, alignment=PP_ALIGN.LEFT)
    # Título
    add_text(slide, Inches(0.20), Inches(0.15), Inches(8.50), Inches(1.0),
             titulo, font_size=28, color=PRETO, bold=True)
    if subtitulo:
        add_text(slide, Inches(0.20), Inches(0.85), Inches(8.50), Inches(0.4),
                 subtitulo, font_size=14, color=RGBColor(0x66, 0x66, 0x66))

def card(slide, left, top, width, height, title, body_lines, title_color=VERDE_IF, bg_color=CINZA_CLARO):
    add_shape(slide, left, top, width, height, bg_color, corner_radius=0.05)
    add_text(slide, left + Inches(0.15), top + Inches(0.08), width - Inches(0.3), Inches(0.35),
             title, font_size=13, color=title_color, bold=True)
    add_bullet_list(slide, left + Inches(0.15), top + Inches(0.42), width - Inches(0.3),
                    height - Inches(0.5), body_lines, font_size=11, color=CINZA_TEXTO)

# =====================================================
# SLIDE 1 — CAPA (padrão PPGESP/IFBA)
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
set_slide_bg(slide, BRANCO)

# Imagem de fundo (arco cinza decorativo)
slide.shapes.add_picture(BG_CAPA, Inches(0), Inches(0), Inches(10), Inches(7.5))

# Logo PPGESP no topo esquerdo
slide.shapes.add_picture(LOGO_PPGESP, Inches(0), Inches(0), Inches(5.40), Inches(1.07))

# Logo cubo no canto superior direito
slide.shapes.add_picture(LOGO_CUBO, Inches(9.02), Inches(0.21), Inches(0.70), Inches(0.80))

# Título do trabalho (centralizado verticalmente)
add_text(slide, Inches(0), Inches(2.70), Inches(9.07), Inches(1.40),
         "Trabalhos Relacionados e Estado da Arte",
         font_size=28, color=PRETO, bold=True, alignment=PP_ALIGN.LEFT)

add_text(slide, Inches(0), Inches(3.80), Inches(9.07), Inches(0.80),
         "Autenticação Adaptativa com Inteligência Artificial,\nBiometria Comportamental e Blockchain",
         font_size=16, color=CINZA_TEXTO, alignment=PP_ALIGN.LEFT)

# Informações do aluno (parte inferior)
add_text(slide, Inches(0), Inches(5.10), Inches(9.07), Inches(0.40),
         "João Guedes de Brito | Mestrando | joaoguedesdebrito@gmail.com",
         font_size=14, color=PRETO, alignment=PP_ALIGN.LEFT)

add_text(slide, Inches(0), Inches(5.50), Inches(9.07), Inches(0.40),
         "Orientador: Prof. Cleber Jorge Lira de Santana",
         font_size=13, color=CINZA_TEXTO, alignment=PP_ALIGN.LEFT)

# Barra inferior cinza escuro
add_shape(slide, Inches(0), Inches(7.09), Inches(10), Inches(0.31), CINZA_BARRA)
add_text(slide, Inches(0.3), Inches(7.10), Inches(7), Inches(0.25),
         "IFBA – PPGESP – Abril 2026",
         font_size=8, color=BRANCO, alignment=PP_ALIGN.LEFT)

# =====================================================
# SLIDE 2 — AGENDA
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Agenda")

items = [
    ("01", "Contexto e Objetivo da Revisão"),
    ("02", "Quadro Comparativo Geral"),
    ("03", "Pesquisa no INPI"),
    ("04", "Trabalhos Acadêmicos (12 estudos)"),
    ("05", "Soluções da Indústria"),
    ("06", "Síntese Comparativa e Diferenciais"),
    ("07", "Lacunas Identificadas na Literatura"),
]
for i, (num, texto) in enumerate(items):
    y = Inches(1.5) + Inches(i * 0.7)
    add_shape(slide, Inches(0.8), y, Inches(0.55), Inches(0.45), VERDE_IF, corner_radius=0.1)
    add_text(slide, Inches(0.82), y + Inches(0.04), Inches(0.50), Inches(0.35),
             num, font_size=16, color=BRANCO, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.6), y + Inches(0.05), Inches(7), Inches(0.35),
             texto, font_size=16, color=CINZA_TEXTO)

# =====================================================
# SLIDE 3 — CONTEXTO
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Contexto e Objetivo da Revisão")

add_text(slide, Inches(0.3), Inches(1.3), Inches(9.3), Inches(1.0),
         "O sistema proposto integra múltiplas tecnologias em uma plataforma única de autenticação inteligente. "
         "Esta revisão identifica e compara trabalhos da literatura acadêmica e da indústria que abordam, "
         "de forma isolada ou parcial, os mesmos eixos temáticos.",
         font_size=14, color=CINZA_TEXTO)

card(slide, Inches(0.3), Inches(2.6), Inches(2.9), Inches(3.8),
     "Machine Learning",
     ["Isolation Forest (40%)", "Random Forest (30%)", "Deep Learning (30%)", "Score unificado 0.0-1.0"],
     title_color=VERDE_IF)

card(slide, Inches(3.5), Inches(2.6), Inches(2.9), Inches(3.8),
     "Biometria e Contexto",
     ["Dinâmica de digitação", "Geolocalização por IP", "Análise de sessão", "Deriva comportamental"],
     title_color=VERDE_IF)

card(slide, Inches(6.7), Inches(2.6), Inches(2.9), Inches(3.8),
     "Blockchain e XAI",
     ["Hyperledger Fabric", "Ethereum (Web3j)", "Auditoria imutável", "Explicabilidade via SHAP"],
     title_color=VERDE_IF)

# =====================================================
# SLIDE 4 — QUADRO COMPARATIVO (parte 1)
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Quadro Comparativo Geral", "Trabalhos Acadêmicos (1/2)")

rows = 8
cols = 5
tbl = slide.shapes.add_table(rows, cols, Inches(0.2), Inches(1.4), Inches(9.6), Inches(5.3)).table
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(0.7)
tbl.columns[2].width = Inches(2.4)
tbl.columns[3].width = Inches(2.1)
tbl.columns[4].width = Inches(1.8)

headers = ["Trabalho", "Ano", "Objetivo", "Experimento", "Métricas"]
data = [
    ["Adaptive Biometric Auth. Systems (ScienceDirect)", "2023", "Mapear estado da arte em autenticação biométrica adaptativa", "Revisão de 80+ trabalhos", "FAR, FRR, EER"],
    ["AI-Driven Adaptive Auth. (ESRG)", "2024", "Reduzir FAR/FRR com ML multimodal", "Testes com biometria multimodal", "FAR -27%, FRR -35%"],
    ["Continuous Smartphone Auth. (MDPI)", "2026", "Autenticação contínua em dispositivos móveis", "Pipeline CNN+LSTM+GRU", "Acurácia, EER, F1"],
    ["Keystroke Dynamics (Springer)", "2025", "Autenticação por padrões de digitação", "KNN, RF, LGBM em texto fixo", "EER 2.67%-6.61%"],
    ["Blockchain Audit Trail (MDPI)", "2021", "Auditoria imutável com blockchain", "Smart contracts", "Throughput, latência"],
    ["Harpocrates (Virginia Tech)", "2022", "Logs imutáveis com privacidade", "Hyperledger Fabric em EC2", "Latência, overhead"],
    ["IAM + Hyperledger + OAuth (ScienceDirect)", "2023", "IAM com blockchain para saúde", "Protótipo Hyperledger", "Escalabilidade, segurança"],
]

for j, h in enumerate(headers):
    cell = tbl.cell(0, j)
    cell.text = h
    for p in cell.text_frame.paragraphs:
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = BRANCO
        p.font.name = "Calibri"
    cell.fill.solid()
    cell.fill.fore_color.rgb = VERDE_ESCURO

for i, row_data in enumerate(data):
    for j, val in enumerate(row_data):
        cell = tbl.cell(i+1, j)
        cell.text = val
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(8)
            p.font.color.rgb = CINZA_TEXTO
            p.font.name = "Calibri"
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRANCO if i % 2 == 0 else CINZA_CLARO

# =====================================================
# SLIDE 5 — QUADRO COMPARATIVO (parte 2)
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Quadro Comparativo Geral", "Trabalhos Acadêmicos (2/2) + Indústria")

rows = 8
cols = 5
tbl = slide.shapes.add_table(rows, cols, Inches(0.2), Inches(1.4), Inches(9.6), Inches(5.3)).table
tbl.columns[0].width = Inches(2.6)
tbl.columns[1].width = Inches(0.7)
tbl.columns[2].width = Inches(2.4)
tbl.columns[3].width = Inches(2.1)
tbl.columns[4].width = Inches(1.8)

headers = ["Trabalho", "Ano", "Objetivo", "Experimento", "Métricas"]
data = [
    ["Risk-Based Auth. com ML (IEEE)", "2018", "Classificar logins por nível de risco", "Modelos ML para scoring", "Acurácia, precisão, recall"],
    ["Trust Engine Behavioral (IJETCSIT)", "2024", "Score de confiança dinâmico passwordless", "Avaliação em móveis/wearables", "TAR>90%, FAR<5%, EER~6.5%"],
    ["AI Continuous Auth. Zero Trust (ResearchGate)", "2025", "Decisões adaptativas com RL", "Framework com Reinforcement Learning", "Redução falsos positivos"],
    ["XAI para IDS (Frontiers in AI)", "2025", "Explicabilidade em sistemas de segurança", "Análise SHAP/LIME em IDS", "Interpretabilidade, fidelidade"],
    ["Adaptive Security + Drift (Al-Qadisiyah)", "2024", "Detecção de deriva comportamental", "LSTM+Autoencoder+KL-divergence", "Acurácia de detecção"],
    ["IBM Trusteer (Indústria)", "2016+", "Prevenção de fraude com biometria IA", "Produção em inst. financeiras", "Proprietário (não publicado)"],
    ["BioCatch (Indústria)", "2011+", "Biometria comportamental cognitiva", "Produção em bancos/fintechs", "Proprietário (não publicado)"],
]

for j, h in enumerate(headers):
    cell = tbl.cell(0, j)
    cell.text = h
    for p in cell.text_frame.paragraphs:
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = BRANCO
        p.font.name = "Calibri"
    cell.fill.solid()
    cell.fill.fore_color.rgb = VERDE_ESCURO

for i, row_data in enumerate(data):
    for j, val in enumerate(row_data):
        cell = tbl.cell(i+1, j)
        cell.text = val
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(8)
            p.font.color.rgb = CINZA_TEXTO
            p.font.name = "Calibri"
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRANCO if i % 2 == 0 else CINZA_CLARO

# =====================================================
# SLIDE 6 — PESQUISA INPI
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Pesquisa no INPI", "Instituto Nacional da Propriedade Industrial")

add_shape(slide, Inches(0.3), Inches(1.5), Inches(4.4), Inches(2.3), CINZA_CLARO, corner_radius=0.05)
add_text(slide, Inches(0.45), Inches(1.55), Inches(4.1), Inches(0.3),
         "Resultados da Busca", font_size=13, color=VERDE_IF, bold=True)
add_bullet_list(slide, Inches(0.45), Inches(1.9), Inches(4.1), Inches(1.7), [
    "47 pedidos de patentes relacionados a blockchain",
    "Nenhum combinando todos os elementos da proposta",
    "Sem registros de biometria comportamental + ML + blockchain",
    "Termos: autenticação adaptativa, login inteligente",
], font_size=11, color=CINZA_TEXTO)

add_shape(slide, Inches(5.0), Inches(1.5), Inches(4.6), Inches(2.3), CINZA_CLARO, corner_radius=0.05)
add_text(slide, Inches(5.15), Inches(1.55), Inches(4.3), Inches(0.3),
         "Diretrizes de IA no INPI (Ago/2025)", font_size=13, color=VERDE_IF, bold=True)
add_bullet_list(slide, Inches(5.15), Inches(1.9), Inches(4.3), Inches(1.7), [
    "Consulta pública sobre patentes de IA",
    "Categorias: modelos de IA, invenções baseadas em IA, invenções assistidas por IA",
    "Amadurecimento regulatório para proteção futura",
], font_size=11, color=CINZA_TEXTO)

add_shape(slide, Inches(0.3), Inches(4.1), Inches(9.3), Inches(2.5), RGBColor(0xE8, 0xF5, 0xE9), corner_radius=0.05)
add_text(slide, Inches(0.45), Inches(4.15), Inches(9.0), Inches(0.3),
         "Relação com a Proposta", font_size=13, color=VERDE_IF, bold=True)
add_text(slide, Inches(0.45), Inches(4.55), Inches(9.0), Inches(1.8),
         "A ausência de patentes no INPI que integrem autenticação adaptativa com ML, biometria comportamental, "
         "blockchain e explicabilidade de IA evidencia o caráter inovador da proposta. O sistema se posiciona "
         "em uma lacuna de propriedade intelectual no cenário brasileiro, combinando elementos que, "
         "individualmente, já possuem registro, mas que em conjunto representam uma contribuição original.",
         font_size=12, color=CINZA_TEXTO)

# =====================================================
# SLIDES 7-12 — TRABALHOS ACADÊMICOS (detalhados)
# =====================================================

trabalhos = [
    {
        "titulo": "Adaptive Biometric Authentication Systems",
        "fonte": "ScienceDirect, 2023 – Revisão Sistemática",
        "descricao": (
            "Revisão sistemática que analisa mais de 80 trabalhos sobre sistemas de autenticação "
            "biométrica adaptativos. Mapeia como esses sistemas ajustam dinamicamente seus processos "
            "de amostragem e reconhecimento em resposta a mudanças no ambiente operacional. "
            "Identifica desafios críticos como ameaças à privacidade, spoofing e acessibilidade."
        ),
        "comparacao": (
            "O sistema proposto implementa concretamente uma solução adaptativa com ensemble de ML "
            "(Isolation Forest + Random Forest + Deep Learning), MFA condicional baseado em risco, "
            "e detecção de deriva comportamental. Adiciona blockchain e explicabilidade de IA — "
            "elementos não cobertos na revisão."
        ),
    },
    {
        "titulo": "AI-Driven Adaptive Auth. for Multi-Modal Biometrics",
        "fonte": "ESRG Journal, 2024",
        "descricao": (
            "Framework com ML para ajuste dinâmico de parâmetros de autenticação baseado em dados "
            "contextuais e comportamentais. Resultados: redução de 27% na FAR e 35% na FRR "
            "comparado a métodos estáticos tradicionais. Utiliza biometria multimodal para perfis adaptativos."
        ),
        "comparacao": (
            "A proposta compartilha o princípio de ajuste dinâmico, mas implementa ensemble de 3 algoritmos "
            "com pesos configuráveis (IF 40%, RF 30%, DL 30%) e regras de decisão escalonadas em 4 níveis "
            "(permitir, MFA, MFA+alerta, bloquear). Integra blockchain para registro imutável de cada decisão."
        ),
    },
    {
        "titulo": "Continuous Smartphone Auth. + Keystroke Dynamics",
        "fonte": "MDPI 2026 / Springer 2025",
        "descricao": (
            "O trabalho da MDPI utiliza pipeline de CNN+LSTM+GRU para autenticação contínua em móveis. "
            "O da Springer aplica KNN, RF e LGBM para keystroke dynamics, alcançando EER de 2.67% a 6.61%. "
            "Ambos demonstram a eficácia de ensembles e biometria comportamental."
        ),
        "comparacao": (
            "O sistema proposto integra dinâmica de digitação como uma das 14 features do ensemble (não isoladamente), "
            "combinando com geolocalização, padrão de sessão e contexto. Adiciona explicabilidade SHAP "
            "e score persistente em blockchain — aspectos não abordados nos trabalhos comparados."
        ),
    },
    {
        "titulo": "Blockchain Audit Trail + Harpocrates",
        "fonte": "MDPI 2021 / Virginia Tech 2022",
        "descricao": (
            "O MDPI propõe mecanismo genérico de audit trail com smart contracts. "
            "Harpocrates implementa logs imutáveis com preservação de privacidade sobre Hyperledger Fabric "
            "em Amazon EC2, combinando técnicas criptográficas com imutabilidade blockchain."
        ),
        "comparacao": (
            "O sistema proposto registra especificamente eventos de autenticação (hash, tipo de evento, "
            "ID hasheado SHA-3, IP, score, decisão, timestamp) em blockchain dual (Hyperledger Fabric + Ethereum). "
            "Opera em modo simulação para dev e redes reais em produção — flexibilidade não presente nos comparados."
        ),
    },
    {
        "titulo": "IAM Hyperledger + OAuth 2.0 / Risk-Based Auth.",
        "fonte": "ScienceDirect 2023 / IEEE 2018",
        "descricao": (
            "O trabalho ScienceDirect combina Hyperledger Fabric com OAuth 2.0 para IAM no setor de saúde. "
            "O IEEE (2018) é um trabalho pioneiro de autenticação baseada em risco com ML, "
            "classificando logins por nível de risco com features contextuais básicas."
        ),
        "comparacao": (
            "O sistema proposto substitui autenticação estática (OAuth) por score dinâmico de IA, "
            "com 4 níveis de resposta em vez de decisão binária. Expande o trabalho de 2018 "
            "com ensemble de 3 algoritmos, 14 features, biometria comportamental, blockchain e explicabilidade."
        ),
    },
    {
        "titulo": "Trust Engine + Zero Trust + XAI + Drift Detection",
        "fonte": "IJETCSIT 2024 / ResearchGate 2025 / Frontiers 2025 / Al-Qadisiyah 2024",
        "descricao": (
            "Trust Engine: score dinâmico passwordless com TAR>90% e FAR<5%. Zero Trust: RL para enforcement "
            "adaptativo. XAI: revisão de SHAP/LIME para IDS. Drift Detection: LSTM+Autoencoder "
            "com KL-divergence para detectar mudanças de comportamento."
        ),
        "comparacao": (
            "O sistema proposto integra TODOS esses elementos em uma única plataforma: "
            "score de confiança (como Trust Engine), decisões adaptativas (como Zero Trust), "
            "explicabilidade SHAP (como XAI), e detecção de deriva (como Drift Detection) — "
            "com o diferencial de registro imutável em blockchain de cada decisão e explicação."
        ),
    },
]

for t in trabalhos:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_padrao(slide, t["titulo"], t["fonte"])

    # Caixa de descrição
    add_shape(slide, Inches(0.2), Inches(1.4), Inches(9.5), Inches(2.2), CINZA_CLARO, corner_radius=0.05)
    add_text(slide, Inches(0.35), Inches(1.45), Inches(9.2), Inches(0.3),
             "Descrição do Trabalho", font_size=12, color=VERDE_IF, bold=True)
    add_text(slide, Inches(0.35), Inches(1.80), Inches(9.2), Inches(1.6),
             t["descricao"], font_size=12, color=CINZA_TEXTO)

    # Caixa de comparação
    add_shape(slide, Inches(0.2), Inches(3.85), Inches(9.5), Inches(2.8), RGBColor(0xE8, 0xF5, 0xE9), corner_radius=0.05)
    add_text(slide, Inches(0.35), Inches(3.90), Inches(9.2), Inches(0.3),
             "Comparação com a Proposta", font_size=12, color=VERDE_IF, bold=True)
    add_text(slide, Inches(0.35), Inches(4.25), Inches(9.2), Inches(2.2),
             t["comparacao"], font_size=12, color=CINZA_TEXTO)

# =====================================================
# SLIDE 13 — SOLUÇÕES DA INDÚSTRIA
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Soluções da Indústria", "IBM Trusteer e BioCatch")

# IBM Card
add_shape(slide, Inches(0.2), Inches(1.4), Inches(4.6), Inches(5.2), CINZA_CLARO, corner_radius=0.05)
add_text(slide, Inches(0.35), Inches(1.45), Inches(4.3), Inches(0.3),
         "IBM Security Trusteer (2016+)", font_size=14, color=VERDE_IF, bold=True)
add_text(slide, Inches(0.35), Inches(1.85), Inches(4.3), Inches(1.3),
         "Suíte de prevenção de fraude com IA, biometria comportamental e "
         "inteligência global. Analisa sinais de identidade digital e "
         "informações de dispositivo. Utilizado por centenas de instituições financeiras.",
         font_size=11, color=CINZA_TEXTO)
add_text(slide, Inches(0.35), Inches(3.3), Inches(4.3), Inches(0.25),
         "vs. Proposta:", font_size=11, color=VERDE_IF, bold=True)
add_text(slide, Inches(0.35), Inches(3.6), Inches(4.3), Inches(2.5),
         "Trusteer é proprietário/black-box. A proposta oferece: explicabilidade SHAP, "
         "blockchain para auditoria, arquitetura aberta e reprodutível, "
         "métricas e decisões transparentes.",
         font_size=11, color=CINZA_TEXTO)

# BioCatch Card
add_shape(slide, Inches(5.1), Inches(1.4), Inches(4.6), Inches(5.2), CINZA_CLARO, corner_radius=0.05)
add_text(slide, Inches(5.25), Inches(1.45), Inches(4.3), Inches(0.3),
         "BioCatch (2011+)", font_size=14, color=VERDE_IF, bold=True)
add_text(slide, Inches(5.25), Inches(1.85), Inches(4.3), Inches(1.3),
         "Biometria comportamental cognitiva e física para prevenção de fraude. "
         "Analisa interações humano-dispositivo com IA. "
         "Diferencia usuários legítimos de fraudadores. Adotado por bancos globais.",
         font_size=11, color=CINZA_TEXTO)
add_text(slide, Inches(5.25), Inches(3.3), Inches(4.3), Inches(0.25),
         "vs. Proposta:", font_size=11, color=VERDE_IF, bold=True)
add_text(slide, Inches(5.25), Inches(3.6), Inches(4.3), Inches(2.5),
         "BioCatch é SaaS proprietário para fraude financeira. A proposta é "
         "voltada para autenticação web geral, com ensemble explícito de 3 algoritmos, "
         "blockchain dual, explicabilidade SHAP e detecção de deriva.",
         font_size=11, color=CINZA_TEXTO)

# =====================================================
# SLIDE 14 — SÍNTESE COMPARATIVA (tabela)
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Síntese Comparativa", "Diferenciais exclusivos da proposta")

rows = 10
cols = 4
tbl = slide.shapes.add_table(rows, cols, Inches(0.2), Inches(1.4), Inches(9.6), Inches(5.3)).table
tbl.columns[0].width = Inches(3.8)
tbl.columns[1].width = Inches(1.6)
tbl.columns[2].width = Inches(2.1)
tbl.columns[3].width = Inches(2.1)

headers = ["Característica", "Proposta", "Literatura", "Indústria"]
data = [
    ["Ensemble de 3 algoritmos (IF+RF+DL) com pesos configuráveis", "SIM", "Parcial (1-2 modelos)", "Não divulgado"],
    ["Score de confiança com 4 níveis de decisão", "SIM", "Binário ou ternário", "Não divulgado"],
    ["Blockchain dual (Hyperledger Fabric + Ethereum)", "SIM", "Uma rede ou nenhuma", "Não utilizado"],
    ["Explicabilidade via SHAP para autenticação", "SIM", "Aplicado em IDS", "Não disponível"],
    ["Detecção de deriva comportamental integrada", "SIM", "Trabalhos isolados", "Parcial"],
    ["Dinâmica de digitação + contexto + geolocalização", "SIM", "Geralmente unimodal", "Multimodal (proprietário)"],
    ["Trilha de auditoria imutável das decisões de IA", "SIM", "Blockchain genérico", "Não disponível"],
    ["Código aberto e reprodutível", "SIM", "Parcial", "Não"],
    ["MFA adaptativo condicionado ao score", "SIM", "Raro", "Parcial"],
]

for j, h in enumerate(headers):
    cell = tbl.cell(0, j)
    cell.text = h
    for p in cell.text_frame.paragraphs:
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = BRANCO
        p.font.name = "Calibri"
    cell.fill.solid()
    cell.fill.fore_color.rgb = VERDE_ESCURO

for i, row_data in enumerate(data):
    for j, val in enumerate(row_data):
        cell = tbl.cell(i+1, j)
        cell.text = val
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(9)
            p.font.name = "Calibri"
            if j == 1 and val == "SIM":
                p.font.color.rgb = VERDE_IF
                p.font.bold = True
            else:
                p.font.color.rgb = CINZA_TEXTO
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRANCO if i % 2 == 0 else CINZA_CLARO

# =====================================================
# SLIDE 15 — LACUNAS NA LITERATURA
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Lacunas Identificadas na Literatura", "Contribuições originais da proposta")

lacunas = [
    ("01", "Explicabilidade + Autenticação + Blockchain",
     "Nenhum trabalho revisado implementa explicabilidade de IA especificamente para decisões de autenticação com registro em blockchain."),
    ("02", "Blockchain Dual para Auditoria",
     "A combinação de Hyperledger Fabric (permissionada) e Ethereum (interoperável) para auditoria de eventos de autenticação não foi encontrada."),
    ("03", "Ensemble de 3 Paradigmas de ML",
     "A combinação de detecção de anomalias (IF), classificação (RF) e aprendizado profundo (DL) com pesos configuráveis para autenticação é contribuição original."),
    ("04", "Resposta Escalonada em 4 Níveis",
     "A granularidade de resposta (permitir, MFA, MFA+alerta, bloquear) baseada em score contínuo supera as abordagens binárias predominantes."),
]

colors_lacuna = [VERDE_IF, VERDE_ESCURO, LARANJA, VERMELHO_IF]

for i, (num, titulo, desc) in enumerate(lacunas):
    y = Inches(1.5) + Inches(i * 1.3)
    add_shape(slide, Inches(0.3), y, Inches(0.55), Inches(0.45), colors_lacuna[i], corner_radius=0.1)
    add_text(slide, Inches(0.32), y + Inches(0.04), Inches(0.50), Inches(0.35),
             num, font_size=16, color=BRANCO, bold=True, alignment=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.1), y, Inches(8.5), Inches(0.35),
             titulo, font_size=14, color=PRETO, bold=True)
    add_text(slide, Inches(1.1), y + Inches(0.38), Inches(8.5), Inches(0.8),
             desc, font_size=11, color=CINZA_TEXTO)

# =====================================================
# SLIDE 16 — REFERÊNCIAS
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide_padrao(slide, "Referências Bibliográficas")

refs = [
    "[1] Adaptive Biometric Auth. Systems — ScienceDirect, 2023",
    "[2] AI-Driven Adaptive Auth. Multi-Modal — ESRG Journal, 2024",
    "[3] Continuous Smartphone Auth. — MDPI Mathematics, 2026",
    "[4] Keystroke Dynamics — Springer Discover Applied Sciences, 2025",
    "[5] Blockchain Audit Trail — MDPI Algorithms, 2021",
    "[6] Harpocrates — arXiv / Virginia Tech, 2022",
    "[7] IAM + Hyperledger + OAuth 2.0 — ScienceDirect, 2023",
    "[8] Risk-Based Auth. com ML — IEEE, 2018",
    "[9] Trust Engine Behavioral — IJETCSIT, 2024",
    "[10] AI Continuous Auth. Zero Trust — ResearchGate, 2025",
    "[11] XAI para IDS — Frontiers in AI, 2025",
    "[12] Adaptive Security + Drift — Al-Qadisiyah Journal, 2024",
    "[13] IBM Security Trusteer — ibm.com",
    "[14] BioCatch Behavioral Biometrics — biocatch.com",
    "[15] INPI Diretrizes de IA — gov.br/inpi, 2025",
]

add_bullet_list(slide, Inches(0.3), Inches(1.3), Inches(9.3), Inches(5.5),
                refs, font_size=10, color=CINZA_TEXTO, bullet_char="")

# =====================================================
# SLIDE 17 — ENCERRAMENTO
# =====================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, BRANCO)

# Imagem de fundo decorativa
slide.shapes.add_picture(BG_CAPA, Inches(0), Inches(0), Inches(10), Inches(7.5))

# Logo PPGESP
slide.shapes.add_picture(LOGO_PPGESP, Inches(0), Inches(0), Inches(5.40), Inches(1.07))

# Logo cubo
slide.shapes.add_picture(LOGO_CUBO, Inches(9.02), Inches(0.21), Inches(0.70), Inches(0.80))

add_text(slide, Inches(0), Inches(2.8), Inches(10), Inches(1),
         "Obrigado!",
         font_size=44, color=PRETO, bold=True, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(0), Inches(3.8), Inches(10), Inches(0.5),
         "Sistema Autônomo de Autenticação Inteligente",
         font_size=16, color=CINZA_TEXTO, bold=True, alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(0), Inches(4.3), Inches(10), Inches(0.5),
         "Trabalhos Relacionados e Estado da Arte",
         font_size=13, color=RGBColor(0x66, 0x66, 0x66), alignment=PP_ALIGN.CENTER)

add_text(slide, Inches(0), Inches(5.2), Inches(10), Inches(0.4),
         "João Guedes de Brito | joaoguedesdebrito@gmail.com",
         font_size=12, color=CINZA_TEXTO, alignment=PP_ALIGN.CENTER)

# Barra inferior
add_shape(slide, Inches(0), Inches(7.09), Inches(10), Inches(0.31), CINZA_BARRA)
add_text(slide, Inches(0.3), Inches(7.10), Inches(7), Inches(0.25),
         "IFBA – PPGESP – Abril 2026",
         font_size=8, color=BRANCO, alignment=PP_ALIGN.LEFT)

# =====================================================
# SALVAR
# =====================================================
output_path = r"C:\dev\projetos\sistema-autonomo\docs\trabalhos-relacionados.pptx"
prs.save(output_path)
print(f"Apresentação salva em: {output_path}")
print(f"Total de slides: {len(prs.slides)}")
