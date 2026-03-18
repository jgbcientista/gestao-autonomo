import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

DOCX_PATH = 'doc-mestrado/doc-qualificacao-ifba.docx'
doc = Document(DOCX_PATH)

# The real style IDs are: Ttulo1 (Heading 1), Ttulo2 (Heading 2), BodyText -> check
print("=== Verificando style IDs reais ===")
for style in doc.styles:
    if style.name in ('Heading 1', 'Heading 2', 'Body Text', 'Normal'):
        print(f"  {style.name} -> style_id={style.style_id}")

# Fix 1: MOTOR H1 (para 289) - change pStyle from Heading1 to Ttulo1
print("\n=== Corrigindo estilo do H1 MOTOR ===")
p289 = doc.paragraphs[289]
pPr = p289._element.find(qn('w:pPr'))
pStyle = pPr.find(qn('w:pStyle'))
old_style = pStyle.get(qn('w:val'))
pStyle.set(qn('w:val'), 'Ttulo1')
print(f"  Para 289: {old_style} → Ttulo1 ({p289.text[:50]})")

# Fix 2: All H2s in MOTOR chapter - change pStyle from Heading2 to Ttulo2
# These should be: Arquitetura, Critérios, Processo, Limiares, Pipeline, Integração
print("\n=== Corrigindo estilos dos H2s MOTOR ===")

# Find MOTOR H1 and DESENVOLVIMENTO H1 to know the range
motor_start = 289
dev_start = None
for i, p in enumerate(doc.paragraphs):
    if p.style and p.style.name == 'Heading 1' and 'DESENVOLVIMENTO' in p.text:
        dev_start = i
        break

print(f"  MOTOR range: {motor_start} to {dev_start}")

h2_titles = [
    'Arquitetura do Ensemble de Modelos',
    'Critérios Comportamentais (14 Dimensões)',
    'Processo de Treinamento com Dados Reais',
    'Limiares de Decisão Autônoma',
    'Pipeline de Inferência em Tempo Real',
    'Integração com Blockchain para Auditoria',
]

fixed_h2 = 0
for i in range(motor_start, dev_start):
    p = doc.paragraphs[i]
    if p.text.strip() in h2_titles:
        pPr = p._element.find(qn('w:pPr'))
        if pPr is not None:
            pStyle = pPr.find(qn('w:pStyle'))
            if pStyle is not None:
                old = pStyle.get(qn('w:val'))
                pStyle.set(qn('w:val'), 'Ttulo2')

                # Also add numPr if missing
                numPr = pPr.find(qn('w:numPr'))
                if numPr is None:
                    numPr = OxmlElement('w:numPr')
                    ilvl = OxmlElement('w:ilvl')
                    ilvl.set(qn('w:val'), '1')
                    numPr.append(ilvl)
                    numId = OxmlElement('w:numId')
                    numId.set(qn('w:val'), '38')
                    numPr.append(numId)
                    pPr.append(numPr)

                print(f"  {i}: {old} → Ttulo2 + numId=38 ({p.text[:50]})")
                fixed_h2 += 1

print(f"  Total H2s corrigidos: {fixed_h2}")

# Fix 3: Body text paragraphs - change BodyText to correct style
print("\n=== Corrigindo estilo Body Text ===")
# Check what the real Body Text style_id is
bt_style_id = None
for style in doc.styles:
    if style.name == 'Body Text':
        bt_style_id = style.style_id
        break
print(f"  Body Text style_id: {bt_style_id}")

fixed_bt = 0
for i in range(motor_start, dev_start):
    p = doc.paragraphs[i]
    pPr = p._element.find(qn('w:pPr'))
    if pPr is not None:
        pStyle = pPr.find(qn('w:pStyle'))
        if pStyle is not None:
            val = pStyle.get(qn('w:val'))
            if val == 'BodyText' and bt_style_id and bt_style_id != 'BodyText':
                pStyle.set(qn('w:val'), bt_style_id)
                fixed_bt += 1
print(f"  {fixed_bt} parágrafos Body Text corrigidos")

doc.save(DOCX_PATH)

# Verify
print("\n=== Verificação ===")
doc2 = Document(DOCX_PATH)

print("\nCapítulos:")
for i, p in enumerate(doc2.paragraphs):
    if p.style and p.style.name == 'Heading 1' and p.text.strip() and i > 180:
        print(f"  {i}: {p.text[:70]}")

print("\nSubcapítulos MOTOR:")
for i, p in enumerate(doc2.paragraphs):
    if p.style and p.style.name == 'Heading 2' and p.text.strip() and i > 289 and i < dev_start + 79:
        pPr = p._element.find(qn('w:pPr'))
        numPr = pPr.find(qn('w:numPr')) if pPr is not None else None
        nid = ''
        if numPr is not None:
            nid_el = numPr.find(qn('w:numId'))
            nid = nid_el.get(qn('w:val')) if nid_el is not None else '?'
        print(f"  {i}: {p.text[:50]} (numId={nid})")

print("\nBoxes:")
for i, p in enumerate(doc2.paragraphs):
    xml_str = etree.tostring(p._element, pretty_print=True).decode()
    if 'w:drawing' in xml_str:
        texts = [t.text for t in p._element.iter(qn('w:t')) if t.text and t.text.strip()]
        combined = ' '.join(texts)
        if 'CAPÍTULO' in combined:
            for j in range(i+1, min(i+15, len(doc2.paragraphs))):
                if doc2.paragraphs[j].style and doc2.paragraphs[j].style.name == 'Heading 1' and doc2.paragraphs[j].text.strip():
                    print(f"  {combined} → {doc2.paragraphs[j].text[:50]}")
                    break

print(f"\nSalvo: {DOCX_PATH}")
