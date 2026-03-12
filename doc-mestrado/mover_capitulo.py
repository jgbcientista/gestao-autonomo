import sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

doc = Document('doc-qualificacao-ifba-12-03-2026.docx')
body = doc.element.body
numbering = doc.part.numbering_part.numbering_definitions._numbering

# === STEP 1: Move MOTOR chapter (498-545) to before RESULTADOS PARCIAIS (447) ===
print("=== Movendo capítulo MOTOR DE IA para antes de RESULTADOS PARCIAIS ===")

# Collect XML elements for MOTOR chapter (498 to 545 inclusive)
motor_elements = []
for i in range(498, 546):
    motor_elements.append(doc.paragraphs[i]._element)

# Get the RESULTADOS PARCIAIS element as insertion reference
resultados_elem = doc.paragraphs[447]._element

# Move each element before RESULTADOS PARCIAIS
for elem in motor_elements:
    resultados_elem.addprevious(elem)

print(f"  Movidos {len(motor_elements)} parágrafos para antes de RESULTADOS PARCIAIS")

# === STEP 2: Update numbering - MOTOR becomes chapter 4, RESULTADOS becomes 5 ===
print("\n=== Atualizando numeração dos capítulos ===")

# Current numbering:
# Chapter 4 (RESULTADOS): numId=4, abstractNumId=36, start=4
# Chapter 5 (MOTOR): numId=38, abstractNumId=37, start=5
# Chapter 6 (CRONOGRAMA): numId=3, abstractNumId=15, start=6

# New order:
# Chapter 4 (MOTOR): needs start=4
# Chapter 5 (RESULTADOS): needs start=5
# Chapter 6 (CRONOGRAMA): stays start=6

# Update MOTOR's abstractNum (37) start from 5 to 4
for abstract in numbering.findall(qn('w:abstractNum')):
    aid = abstract.get(qn('w:abstractNumId'))
    if aid == '37':  # MOTOR
        for lvl in abstract.findall(qn('w:lvl')):
            if lvl.get(qn('w:ilvl')) == '0':
                start = lvl.find(qn('w:start'))
                if start is not None:
                    start.set(qn('w:val'), '4')
                    print(f"  MOTOR (abstractNum 37): start 5 -> 4")
    elif aid == '36':  # RESULTADOS
        for lvl in abstract.findall(qn('w:lvl')):
            if lvl.get(qn('w:ilvl')) == '0':
                start = lvl.find(qn('w:start'))
                if start is not None:
                    start.set(qn('w:val'), '5')
                    print(f"  RESULTADOS (abstractNum 36): start 4 -> 5")

# === STEP 3: Update TOC entries order ===
print("\n=== Atualizando sumário ===")

for sdt in body.findall(qn('w:sdt')):
    sdtPr = sdt.find(qn('w:sdtPr'))
    if sdtPr is None:
        continue
    docPartObj = sdtPr.find(qn('w:docPartObj'))
    if docPartObj is None:
        continue
    gallery = docPartObj.find(qn('w:docPartGallery'))
    if gallery is None or 'Table of Contents' not in gallery.get(qn('w:val'), ''):
        continue

    sdtContent = sdt.find(qn('w:sdtContent'))
    paras = list(sdtContent.findall(qn('w:p')))

    # Find MOTOR entries (27-33) and RESULTADOS entries (21-26)
    # Current order: ..., RESULTADOS(21-26), MOTOR(27-33), CRONOGRAMA(34-35), ...
    # Need: ..., MOTOR entries, RESULTADOS entries, CRONOGRAMA, ...

    # Find indices by content
    motor_toc_start = None
    motor_toc_end = None
    resultados_toc_start = None
    resultados_toc_end = None

    for idx, p in enumerate(paras):
        text = ''.join(t.text for t in p.iter(qn('w:t')) if t.text)
        if 'RESULTADOS PARCIAIS' in text:
            resultados_toc_start = idx
        if 'MOTOR DE INTELIG' in text:
            motor_toc_start = idx
        if 'CRONOGRAMA' in text:
            if motor_toc_start is not None and motor_toc_end is None:
                motor_toc_end = idx
            if resultados_toc_start is not None and resultados_toc_end is None and motor_toc_start is not None:
                pass  # resultados_toc_end set below

    # RESULTADOS entries: from resultados_toc_start to motor_toc_start-1
    # MOTOR entries: from motor_toc_start to motor_toc_end-1 (before CRONOGRAMA)
    resultados_toc_end = motor_toc_start

    print(f"  RESULTADOS TOC: indices {resultados_toc_start}-{resultados_toc_end-1}")
    print(f"  MOTOR TOC: indices {motor_toc_start}-{motor_toc_end-1}")

    # Collect MOTOR TOC elements
    motor_toc_elems = []
    for idx in range(motor_toc_start, motor_toc_end):
        motor_toc_elems.append(paras[idx])

    # Collect RESULTADOS TOC elements (to insert after MOTOR)
    resultados_toc_elems = []
    for idx in range(resultados_toc_start, resultados_toc_end):
        resultados_toc_elems.append(paras[idx])

    # Move MOTOR entries before RESULTADOS entries
    # = move MOTOR entries before the first RESULTADOS entry
    first_resultados = resultados_toc_elems[0]
    for elem in motor_toc_elems:
        first_resultados.addprevious(elem)

    print(f"  Entradas do sumário reordenadas")

    # Now update the numbering in TOC text
    # MOTOR entries: change 5.x to 4.x
    for elem in motor_toc_elems:
        for t_elem in elem.iter(qn('w:t')):
            if t_elem.text:
                t_elem.text = t_elem.text.replace('5.1', '4.1').replace('5.2', '4.2').replace('5.3', '4.3').replace('5.4', '4.4').replace('5.5', '4.5').replace('5.6', '4.6')

    # RESULTADOS entries: change 4.x to 5.x
    for elem in resultados_toc_elems:
        for t_elem in elem.iter(qn('w:t')):
            if t_elem.text:
                t_elem.text = t_elem.text.replace('4.1', '5.1').replace('4.2', '5.2').replace('4.3', '5.3').replace('4.4', '5.4').replace('4.5', '5.5')

    print(f"  Numeração do sumário atualizada (MOTOR=4.x, RESULTADOS=5.x)")
    break

# === STEP 4: Save ===
doc.save('doc-qualificacao-ifba-12-03-2026.docx')

# === Verify ===
print("\n=== Verificação ===")
doc2 = Document('doc-qualificacao-ifba-12-03-2026.docx')

print("Headings H1:")
for i, p in enumerate(doc2.paragraphs):
    if p.style and p.style.name == 'Heading 1' and p.text.strip() and i > 190:
        print(f"  {i}: {p.text[:70]}")

print("\nSumário:")
for sdt in doc2.element.body.findall(qn('w:sdt')):
    sdtPr = sdt.find(qn('w:sdtPr'))
    if sdtPr is not None:
        docPartObj = sdtPr.find(qn('w:docPartObj'))
        if docPartObj is not None:
            gallery = docPartObj.find(qn('w:docPartGallery'))
            if gallery is not None and 'Table of Contents' in gallery.get(qn('w:val'), ''):
                sdtContent = sdt.find(qn('w:sdtContent'))
                for p in sdtContent.findall(qn('w:p')):
                    text = ''.join(t.text for t in p.iter(qn('w:t')) if t.text)
                    if text.strip():
                        pPr = p.find(qn('w:pPr'))
                        style = ''
                        if pPr is not None:
                            ps = pPr.find(qn('w:pStyle'))
                            if ps is not None:
                                style = ps.get(qn('w:val'), '')
                        print(f"  [{style}] {text[:80]}")
